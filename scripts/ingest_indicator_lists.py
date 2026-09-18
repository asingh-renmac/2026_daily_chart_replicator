"""Fold the CFNAI and BBKI published indicator lists into the ratified store.

WHAT THESE ARE
  Two institutions publishing the Haver mnemonics behind their own indexes:
  the Chicago Fed's National Activity Index (85 series, in a PDF table) and the
  Brave-Butters-Kelley indexes (498 rows, sheet 490 of the data dictionary). Each row
  pairs a plain-English description with a Haver ticker, which is exactly the shape of
  learned_descriptors.json -- description in, ticker out.

WHY THEY ARE NOT SIMPLY TRUSTED
  The series book was confirmed ticker-by-ticker against live DLX before anything used it.
  These were not: they are correct for the institution that published them, on the date
  they published, against THEIR Haver licence. A mnemonic here may be discontinued, may be
  in a database we do not take, or may not exist for us at all. So every candidate is read
  back from the catalogue and only a verified one is written. An unverified ticker in the
  ratified store is worse than no entry, because the store is consulted BEFORE search and
  its answers read as human-confirmed.

THE FOOTNOTE PROBLEM, AND WHY THE CATALOGUE SETTLES IT
  The CFNAI table appends footnote markers directly to the mnemonic, with no separator:
  "IPMDG@IP5" is IPMDG@IP with footnote 5, and "KRM257" is KRM25 with footnote 7. Nothing
  in the text distinguishes those digits from a mnemonic that genuinely ends in one --
  IP531@IP, IP54@IP and A0M008@BCI all really do. Rather than guess, each token is expanded
  into candidate readings, longest first, and the CATALOGUE decides which exists. That
  turns an unresolvable parsing ambiguity into a lookup.

FORMULAS ARE RECORDED, NEVER STORED
  70 BBKI rows are arithmetic Haver can evaluate (LETPRIVA/PCU), and 137 use functions from
  the authors' own MATLAB toolbox -- FFGR, BFGR, SPLICE, EXTEND_LAST, FISHERPRICE -- which
  Haver has never heard of. cbd (davidakelley/cbd, Kelley being the K in Brave-Butters-
  Kelley) reimplements the 15 standard Haver functions and adds these. Some even reference
  @LOCAL, a private database of theirs. None of it belongs in a store whose contract is
  "this description means this ONE ticker", so both kinds are reported and skipped.

    python scripts/ingest_indicator_lists.py             # dry run, verifies, writes nothing
    python scripts/ingest_indicator_lists.py --apply
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import haver_search as HS  # noqa: E402
import resolve as R  # noqa: E402

PDF = ROOT / "notes" / "cfnai-indicators-list-pdf.pdf"
XLSX = ROOT / "notes" / "BBKI_data_dictionary.xlsx"

# A spec containing any of these is a formula, not a ticker.
FORMULA_CHARS = re.compile(r"[/+*()]|\s-\s")
# CFNAI row: "<description> <0.0nn> <TRANSFORM> <mnemonic+footnotes>"
CFNAI_ROW = re.compile(r"^(.*?)\s+(-?0\.\d+)\s+(DLN|DLV|LN|LV)\s+(\S+)\s*$")


def ticker_candidates(raw: str) -> list[str]:
    """Readings of a mnemonic token, longest first, footnote digits peeled off.

    'IPMDG@IP5' -> IPMDG@IP5, IPMDG@IP.   'KRM257' -> KRM257@USECON, KRM25@USECON, ...
    Bare mnemonics take USECON, which both documents state as their default.
    """
    tok = raw.strip().rstrip(".;")
    code, _, db = tok.partition("@")
    out: list[str] = []
    # Peel DIGITS AND COMMAS, because a row can carry several footnotes at once --
    # "CPG4,6" is CPG with notes 4 and 6, and stopping at the comma never reaches the
    # mnemonic. No Haver mnemonic contains a comma, so peeling one is always right.
    if db:
        seen = db                       # footnotes land after the database: IP5 -> IP
        while seen:
            out.append(f"{code}@{seen}".lower())
            if not (seen[-1].isdigit() or seen[-1] == ","):
                break
            seen = seen[:-1]
    else:
        seen = code                     # bare mnemonic: KRM257 -> KRM25 -> KRM2 -> KRM
        while seen:
            out.append(f"{seen}@usecon".lower())
            if not (seen[-1].isdigit() or seen[-1] == ","):
                break
            seen = seen[:-1]
    return [c for c in out if not c.startswith("@") and "," not in c.split("@")[0]]


def parse_cfnai() -> list[dict]:
    import pdfplumber

    rows: list[dict] = []
    with pdfplumber.open(PDF) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").splitlines():
                m = CFNAI_ROW.match(line.strip())
                if not m:
                    continue
                desc, _w, xf, tok = m.groups()
                desc = desc.replace("(constructed)", "").strip()
                rows.append({"src": "CFNAI", "desc": desc, "raw": tok, "xf": xf})
    return rows


def parse_bbki() -> list[dict]:
    import openpyxl

    ws = openpyxl.load_workbook(XLSX, data_only=True)["490"]
    rows = []
    for r in range(2, ws.max_row + 1):
        desc = (ws.cell(r, 2).value or "").strip()
        spec = str(ws.cell(r, 7).value or "").strip()
        if desc and spec:
            rows.append({"src": "BBKI", "desc": desc, "raw": spec,
                         "xf": str(ws.cell(r, 6).value or "").strip()})
    return rows


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()

    rows = parse_cfnai() + parse_bbki()
    if args.limit:
        rows = rows[:args.limit]
    by_src = Counter(r["src"] for r in rows)
    print(f"parsed: {len(rows)} rows  ({dict(by_src)})\n")

    stats = Counter()
    verified: list[dict] = []
    unverified: list[dict] = []

    for i, e in enumerate(rows, 1):
        if FORMULA_CHARS.search(e["raw"]):
            stats["skipped (formula, not one ticker)"] += 1
            continue
        hit = None
        for cand in ticker_candidates(e["raw"]):
            try:
                meta = HS.get_meta(cand)
            except Exception:
                meta = None
            if meta:
                hit = (cand, meta)
                break
        if hit:
            code, meta = hit
            verified.append({**e, "ticker": code, "meta": meta})
            stats["verified against the catalogue"] += 1
        else:
            unverified.append(e)
            stats["NOT in our catalogue"] += 1
        if i % 50 == 0:
            print(f"  {i:4}/{len(rows)} checked", flush=True)

    print()
    for k, v in sorted(stats.items()):
        print(f"  {k:36} {v}")

    # Write only what verified, and never over a disagreement.
    store_path = R.CLARIFIED_DIR / R._LEARNED_FILE
    store = json.loads(store_path.read_text(encoding="utf-8"))
    before = len(store)
    now = datetime.now(timezone.utc).isoformat()
    added = conflicts = agreed = 0
    conflict_rows: list[str] = []

    # Descriptions naming MORE THAN ONE series, computed across the whole input before a
    # single write. Two BBKI rows normalise to "ip: finished processing" and point at
    # ip564@ip (production) and cu564@ip (capacity utilisation) -- different series whose
    # published descriptions collapse together. Whichever row was read first would win, be
    # wrong half the time, and look deliberate, because a store hit reads as a ratified
    # human answer. learned_lookup already refuses an ambiguous key on READ; a writer that
    # does not honour the same rule never lets that refusal fire.
    owners: dict[str, set] = {}
    for e in verified:
        k = R._norm_key(e["desc"])
        if k:
            owners.setdefault(k, set()).add(e["ticker"])
    ambiguous = {k for k, v in owners.items() if len(v) > 1}
    if ambiguous:
        print(f"\n  {len(ambiguous)} description(s) name two different series -- skipped:")
        for k in sorted(ambiguous)[:6]:
            print(f"     {k[:50]!r} -> {sorted(owners[k])}")

    for e in verified:
        key = R._norm_key(e["desc"])
        if not key or key in ambiguous:
            if key in ambiguous:
                stats["skipped (description names two series)"] += 1
            continue
        cur = store.get(key)
        if isinstance(cur, dict) and cur.get("code"):
            if cur["code"].strip().lower() == e["ticker"]:
                agreed += 1
            else:
                conflicts += 1
                conflict_rows.append(f"{key[:46]!r}\n       store={cur['code']} "
                                     f"vs {e['src']}={e['ticker']}")
            continue
        store[key] = {"code": e["ticker"],
                      "descriptor": e["meta"].get("descriptor", ""),
                      "adjustment": R._sa_norm(e["meta"].get("sa_status") or ""),
                      "added": now,
                      "source": f"indicator_list:{e['src']}"}
        added += 1

    print(f"\n  new entries   : {added}")
    print(f"  already agreed: {agreed}")
    print(f"  CONFLICTS     : {conflicts}   (left exactly as they were)")
    for c in conflict_rows[:8]:
        print(f"     {c}")
    print(f"\n  store: {before} -> {len(store)}")

    if unverified:
        print(f"\n  {len(unverified)} not found in our catalogue (first 12):")
        for e in unverified[:12]:
            print(f"     {e['src']} {e['raw'][:22]:22} {e['desc'][:44]}")

    if not args.apply:
        print("\n  DRY RUN -- nothing written. Re-run with --apply.")
        return 0

    backup = store_path.with_suffix(f".bak-{datetime.now():%Y%m%d-%H%M%S}.json")
    shutil.copy2(store_path, backup)
    store_path.write_text(json.dumps(store, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\n  backup : {backup.name}\n  written: {store_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
